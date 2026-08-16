r"""Provenance-preserving extraction (spec §22–§24).

Every parser returns a list of Segment dicts:
    {"page": int|None, "locator": str, "text": str, "level": str}
so answers can cite page / sheet / slide / cell-range, not just filenames.

Parsers are adapted from the battle-tested V1 code (rag.py) and extended:
PDF (text layer -> OCR -> Docling), DOCX (headings/tables), XLSX (sheet +
range locators), PPTX (slide + notes), TXT/MD/CSV/JSON, EML, images (OCR).
CAD (.dwg/.dxf) is never parsed directly — only derivatives (§24); ezdxf is
used read-only for DXF when available.

All heavy imports are lazy so the package imports without the optional deps;
a missing dependency surfaces as an ExtractionError recorded in the manifest,
never a crash of the whole run (§46).
"""

from __future__ import annotations

import json
import os
from pathlib import Path


class ExtractionError(RuntimeError):
    pass


def seg(text: str, page: int | None = None, locator: str = "",
        level: str = "paragraph") -> dict:
    return {"page": page, "locator": locator, "text": text, "level": level}


# ---------------------------------------------------------------- PDF
def extract_pdf(path: str, ocr: bool = False, ocr_lang: str = "eng+msa") -> list[dict]:
    try:
        import fitz  # PyMuPDF
    except ImportError as e:
        raise ExtractionError(f"PyMuPDF not installed: {e}")
    out = []
    doc = fitz.open(path)
    try:
        for i, page in enumerate(doc):
            t = page.get_text("text")
            if t and t.strip():
                out.append(seg(t, page=i + 1, locator=f"p.{i+1}"))
            elif ocr:
                t = _ocr_pdf_page(page, ocr_lang)
                if t.strip():
                    out.append(seg(t, page=i + 1, locator=f"p.{i+1} (OCR)"))
    finally:
        doc.close()
    return out


def _ocr_pdf_page(page, lang: str) -> str:
    try:
        import io
        import pytesseract
        from PIL import Image
    except ImportError:
        return ""
    try:
        pix = page.get_pixmap(dpi=300)
        img = Image.open(io.BytesIO(pix.tobytes("png")))
    except Exception:
        return ""
    for lg in (lang, "eng"):
        try:
            return pytesseract.image_to_string(img, lang=lg)
        except Exception:
            continue
    return ""


def extract_pdf_docling(path: str) -> list[dict]:
    """Structured parse via Docling (layout + tables + OCR). GPU if available.
    Used by the DEEP-quality ingestion pass; plain extract_pdf is the fast path."""
    try:
        from docling.document_converter import DocumentConverter
    except ImportError as e:
        raise ExtractionError(f"docling not installed: {e}")
    from collections import defaultdict
    res = DocumentConverter().convert(path)
    pages = defaultdict(list)
    for item, _lvl in res.document.iterate_items():
        txt = getattr(item, "text", None)
        if not txt or not txt.strip():
            continue
        pno = 0
        prov = getattr(item, "prov", None)
        if prov:
            pno = getattr(prov[0], "page_no", 0)
        pages[pno].append(txt)
    return [seg("\n".join(t), page=p or None, locator=f"p.{p}" if p else "")
            for p, t in sorted(pages.items())]


# ---------------------------------------------------------------- Office
def extract_docx(path: str) -> list[dict]:
    try:
        import docx
    except ImportError as e:
        raise ExtractionError(f"python-docx not installed: {e}")
    d = docx.Document(path)
    out, section = [], "Document"
    buf: list[str] = []

    def flush():
        if buf:
            out.append(seg("\n".join(buf), locator=section, level="section"))
            buf.clear()

    for p in d.paragraphs:
        if not p.text or not p.text.strip():
            continue
        style = (p.style.name or "").lower() if p.style else ""
        if style.startswith("heading"):
            flush()
            section = p.text.strip()[:120]
            buf.append(p.text.strip())
        else:
            buf.append(p.text)
    flush()
    for ti, tbl in enumerate(d.tables):
        rows = []
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            out.append(seg("\n".join(rows), locator=f"Table {ti+1}", level="row"))
    return out


def extract_xlsx(path: str, max_rows_per_block: int = 60) -> list[dict]:
    try:
        import openpyxl
    except ImportError as e:
        raise ExtractionError(f"openpyxl not installed: {e}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    try:
        for si, ws in enumerate(wb.worksheets):
            block, start_row = [], 1
            for ri, row in enumerate(ws.iter_rows(values_only=True), 1):
                vals = [str(c) for c in row if c is not None]
                if vals:
                    block.append(" | ".join(vals))
                if len(block) >= max_rows_per_block:
                    out.append(seg("\n".join(block), page=si + 1,
                                   locator=f"Sheet '{ws.title}' rows {start_row}-{ri}",
                                   level="row"))
                    block, start_row = [], ri + 1
            if block:
                out.append(seg("\n".join(block), page=si + 1,
                               locator=f"Sheet '{ws.title}' rows {start_row}+",
                               level="row"))
    finally:
        wb.close()
    return out


def extract_pptx(path: str) -> list[dict]:
    try:
        from pptx import Presentation
    except ImportError as e:
        raise ExtractionError(f"python-pptx not installed: {e}")
    prs = Presentation(path)
    out = []
    for si, slide in enumerate(prs.slides, 1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text
                if t and t.strip():
                    parts.append(t.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        try:
            notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        except Exception:
            notes = ""
        if notes.strip():
            parts.append(f"[Notes] {notes.strip()}")
        if parts:
            out.append(seg("\n".join(parts), page=si, locator=f"Slide {si}"))
    return out


# ---------------------------------------------------------------- plain / data
def extract_text(path: str) -> list[dict]:
    raw = Path(path).read_bytes()
    return [seg(raw.decode("utf-8", errors="replace"))]


def extract_json(path: str, max_chars: int = 200_000) -> list[dict]:
    raw = Path(path).read_text(encoding="utf-8", errors="replace")[:max_chars]
    try:
        pretty = json.dumps(json.loads(raw), indent=1, ensure_ascii=False)[:max_chars]
    except Exception:
        pretty = raw
    return [seg(pretty)]


def extract_eml(path: str) -> list[dict]:
    import email
    import email.policy
    with open(path, "rb") as f:
        msg = email.message_from_binary_file(f, policy=email.policy.default)
    header = "\n".join(f"{k}: {msg.get(k, '')}" for k in
                       ("Subject", "From", "To", "Cc", "Date") if msg.get(k))
    body = ""
    try:
        part = msg.get_body(preferencelist=("plain", "html"))
        if part:
            body = part.get_content()
    except Exception:
        pass
    locator = f"{msg.get('Subject', '(no subject)')} — {msg.get('Date', '')}"[:200]
    return [seg(f"{header}\n\n{body}".strip(), locator=locator, level="section")]


def extract_image_ocr(path: str, ocr_lang: str = "eng+msa") -> list[dict]:
    try:
        import pytesseract
        from PIL import Image
    except ImportError as e:
        raise ExtractionError(f"OCR deps not installed: {e}")
    img = Image.open(path)
    for lg in (ocr_lang, "eng"):
        try:
            t = pytesseract.image_to_string(img, lang=lg)
            return [seg(t, locator="image OCR")] if t.strip() else []
        except Exception:
            continue
    return []


# ---------------------------------------------------------------- CAD (read-only derivatives)
def extract_dxf(path: str) -> list[dict]:
    """READ-ONLY text/metadata harvest from a DXF: layer names, block names,
    text entities. Never saves anything back (§24). DWG must first be exported
    to DXF/PDF by an external tool; that derivative is what gets ingested."""
    try:
        import ezdxf
    except ImportError as e:
        raise ExtractionError(f"ezdxf not installed: {e}")
    doc = ezdxf.readfile(path)
    out = []
    layers = [ly.dxf.name for ly in doc.layers]
    if layers:
        out.append(seg("Layers: " + ", ".join(layers), locator="DXF layers",
                       level="section"))
    blocks = [b.name for b in doc.blocks if not b.name.startswith("*")]
    if blocks:
        out.append(seg("Blocks: " + ", ".join(blocks[:500]), locator="DXF blocks",
                       level="section"))
    texts = []
    for e in doc.modelspace():
        if e.dxftype() in ("TEXT", "MTEXT"):
            t = e.plain_text() if hasattr(e, "plain_text") else e.dxf.text
            if t and t.strip():
                texts.append(t.strip())
    if texts:
        out.append(seg("\n".join(texts), locator="DXF text entities"))
    return out


# ---------------------------------------------------------------- dispatch
def extract_any(path: str, ext: str, *, ocr: bool = False,
                ocr_lang: str = "eng+msa", use_docling: bool = False) -> tuple[list[dict], str]:
    """Returns (segments, parser_used)."""
    if ext == ".pdf":
        if use_docling:
            try:
                segs = extract_pdf_docling(path)
                if segs:
                    return segs, "docling"
            except ExtractionError:
                pass  # fall through to fast path
        return extract_pdf(path, ocr=ocr, ocr_lang=ocr_lang), "pymupdf" + ("+ocr" if ocr else "")
    if ext == ".docx":
        return extract_docx(path), "python-docx"
    if ext == ".xlsx":
        return extract_xlsx(path), "openpyxl"
    if ext == ".pptx":
        return extract_pptx(path), "python-pptx"
    if ext in (".txt", ".md", ".csv"):
        return extract_text(path), "text"
    if ext == ".json":
        return extract_json(path), "json"
    if ext == ".eml":
        return extract_eml(path), "eml"
    if ext == ".dxf":
        return extract_dxf(path), "ezdxf-readonly"
    if ext in (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"):
        return extract_image_ocr(path, ocr_lang), "tesseract"
    raise ExtractionError(f"no parser for {ext}")


# ---------------------------------------------------------------- page images (§19/§21)
def render_page_images(path: str, out_dir: Path, file_id: int,
                       max_px: int = 1600, pages: list[int] | None = None) -> list[str]:
    """Render PDF pages to size-capped PNGs under 05_PAGE_IMAGES/<file_id>/.
    V1 lesson baked in: pixel-budget instead of fixed DPI, so A0/A1 tender
    sheets stay ~0.2-0.8MB instead of 18MB."""
    try:
        import fitz
    except ImportError:
        return []
    out_dir = Path(out_dir) / str(file_id)
    out_dir.mkdir(parents=True, exist_ok=True)
    written = []
    doc = fitz.open(path)
    try:
        targets = pages or range(1, len(doc) + 1)
        for pno in targets:
            if pno < 1 or pno > len(doc):
                continue
            out = out_dir / f"p{pno}.png"
            if out.exists():
                written.append(str(out))
                continue
            page = doc[pno - 1]
            maxdim = max(page.rect.width, page.rect.height) or 1.0
            zoom = min(2.0, max_px / maxdim)
            pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom))
            pix.save(str(out))
            written.append(str(out))
    finally:
        doc.close()
    return written
